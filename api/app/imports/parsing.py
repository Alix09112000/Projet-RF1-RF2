"""Lecture et contrôle des fichiers déposés.

Fonctions pures, sans base de données : elles reçoivent le contenu du fichier et le contexte
(mois déjà importés, valeurs du mois précédent) et renvoient erreurs, avertissements et valeurs.
"""

from __future__ import annotations

import csv
import io
import math
import re
import unicodedata
from collections.abc import Collection, Mapping
from dataclasses import dataclass, field
from datetime import date, timedelta
from difflib import get_close_matches
from typing import Any

from openpyxl import load_workbook

from app.imports.spec import (
    CITY,
    COMMERCIAL_HEADERS,
    COMMERCIAL_METRICS,
    COUNTRY_KEY,
    ECHUS_HEADERS,
    ECHUS_SERIES,
    SENIORITY_WARN,
    TOTAL_CHECKS,
    VARIATION_METRICS,
    VARIATION_THRESHOLD,
)

MAX_ISSUES = 100
Row = list[Any]
Previous = Mapping[str, Mapping[date, float]]

MONTH_NAMES = [
    "janvier",
    "février",
    "mars",
    "avril",
    "mai",
    "juin",
    "juillet",
    "août",
    "septembre",
    "octobre",
    "novembre",
    "décembre",
]
_MONTH_PREFIXES = [
    ("janv", 1),
    ("jan", 1),
    ("fev", 2),
    ("mar", 3),
    ("avr", 4),
    ("mai", 5),
    ("juin", 6),
    ("juil", 7),
    ("aou", 8),
    ("sep", 9),
    ("oct", 10),
    ("nov", 11),
    ("dec", 12),
]
_NUMBER = re.compile(r"^-?\d+(?:[.,]\d+)?$")


class FileFormatError(Exception):
    """Fichier illisible : format inconnu, plusieurs feuilles, fichier endommagé."""


@dataclass
class Issue:
    message: str
    line: int | None = None
    column: str | None = None

    def as_dict(self) -> dict[str, Any]:
        return {"message": self.message, "line": self.line, "column": self.column}


@dataclass
class Record:
    city: str
    month: date
    metric: str
    seniority: int
    value: float

    def as_dict(self) -> dict[str, Any]:
        return {
            "city": self.city,
            "month": self.month.isoformat(),
            "metric": self.metric,
            "seniority": self.seniority,
            "value": self.value,
        }


@dataclass
class ParseResult:
    kind: str
    errors: list[Issue] = field(default_factory=list)
    warnings: list[Issue] = field(default_factory=list)
    records: list[Record] = field(default_factory=list)
    months: list[date] = field(default_factory=list)
    rows_read: int = 0
    rows_kept: int = 0
    rows_ignored: int = 0

    @property
    def ok(self) -> bool:
        return not self.errors

    def error(self, message: str, line: int | None = None, column: str | None = None) -> None:
        self.errors.append(Issue(message, line, column))

    def warn(self, message: str, line: int | None = None, column: str | None = None) -> None:
        self.warnings.append(Issue(message, line, column))

    def finish(self) -> ParseResult:
        self.errors = _cap(self.errors, "erreurs")
        self.warnings = _cap(self.warnings, "avertissements")
        return self


def _cap(items: list[Issue], word: str) -> list[Issue]:
    if len(items) <= MAX_ISSUES:
        return items
    return items[:MAX_ISSUES] + [Issue(f"… et {len(items) - MAX_ISSUES} autres {word}.")]


# --- Mise en forme -------------------------------------------------------------------------


def month_label(d: date) -> str:
    return f"{MONTH_NAMES[d.month - 1]} {d.year}"


def fmt_date(d: date) -> str:
    return d.strftime("%d/%m/%Y")


def fmt_num(v: float) -> str:
    if float(v).is_integer():
        return f"{int(v):,}".replace(",", " ")
    return f"{v:,.2f}".replace(",", " ").replace(".", ",")


def plural(n: int, word: str) -> str:
    return f"{n} {word}{'s' if n > 1 else ''}"


def add_months(d: date, n: int) -> date:
    years, month0 = divmod(d.month - 1 + n, 12)
    return date(d.year + years, month0 + 1, 1)


# --- Lecture du fichier --------------------------------------------------------------------


def read_table(content: bytes, filename: str) -> list[Row]:
    """Renvoie les lignes du fichier, en-têtes compris, sans les lignes vides de fin."""
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        rows = _read_xlsx(content)
    elif name.endswith((".csv", ".txt")):
        rows = _read_csv(content)
    elif name.endswith(".xls"):
        raise FileFormatError(
            "Ancien format Excel (.xls) : enregistrez le fichier en .xlsx ou en .csv."
        )
    else:
        raise FileFormatError("Format non reconnu : déposez un fichier .xlsx ou .csv.")
    while rows and _is_empty_row(rows[-1]):
        rows.pop()
    return rows


def _read_xlsx(content: bytes) -> list[Row]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:
        raise FileFormatError("Fichier Excel illisible ou endommagé.") from exc
    try:
        sheets = workbook.worksheets
        if len(sheets) != 1:
            raise FileFormatError(
                f"Le fichier contient {len(sheets)} feuilles : il doit n'en contenir qu'une."
            )
        return [list(row) for row in sheets[0].iter_rows(values_only=True)]
    finally:
        workbook.close()


def _read_csv(content: bytes) -> list[Row]:
    text: str | None = None
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise FileFormatError("Encodage non reconnu : enregistrez le fichier en CSV UTF-8.")
    first_line = text.splitlines()[0] if text else ""
    delimiter = max((";", ",", "\t"), key=first_line.count)
    if first_line.count(delimiter) == 0:
        delimiter = ";"
    return [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


# --- Valeurs ---------------------------------------------------------------------------------


def parse_month(value: Any) -> date | None:
    """Mois au 1er, avec l'année. Renvoie None si illisible ou sans année."""
    if isinstance(value, date):
        return date(value.year, value.month, 1)
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if 20000 <= value <= 80000:  # date Excel stockée en nombre de jours
            d = date(1899, 12, 30) + timedelta(days=int(value))
            return date(d.year, d.month, 1)
        return None
    if not isinstance(value, str):
        return None
    s = value.strip()
    m = re.match(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})\b", s)
    if m:
        return _month(int(m[3]), int(m[2]))
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})\b", s)
    if m:
        return _month(int(m[1]), int(m[2]))
    m = re.match(r"^(\d{1,2})[/.-](\d{4})$", s)
    if m:
        return _month(int(m[2]), int(m[1]))
    m = re.match(r"^(\d{4})[/-](\d{1,2})$", s)
    if m:
        return _month(int(m[1]), int(m[2]))
    return None


def _month(year: int, month: int) -> date | None:
    if 1 <= month <= 12 and 2000 <= year <= 2100:
        return date(year, month, 1)
    return None


def parse_number(value: Any) -> float:
    """Cellule vide = 0. Lève ValueError("formule") ou ValueError("texte")."""
    if value is None:
        return 0.0
    if isinstance(value, bool) or isinstance(value, date):
        raise ValueError("texte")
    if isinstance(value, (int, float)):
        number = float(value)
        return 0.0 if math.isnan(number) else number
    s = str(value).strip()
    if not s:
        return 0.0
    if s.startswith("="):
        raise ValueError("formule")
    cleaned = re.sub(r"[\s\u00a0\u202f]", "", s)
    if _NUMBER.match(cleaned):
        return float(cleaned.replace(",", "."))
    raise ValueError("texte")


def _number_message(line: int, column: str, raw: Any, exc: ValueError) -> str:
    if str(exc) == "formule":
        return (
            f'Ligne {line}, colonne {column} : formule "{_raw(raw)}" au lieu d\'une valeur. '
            "Collez les données en valeurs."
        )
    return f'Ligne {line}, colonne {column} : "{_raw(raw)}" n\'est pas un nombre.'


def _parse_seniority(value: Any) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        n = value
    elif isinstance(value, float) and value.is_integer():
        n = int(value)
    elif isinstance(value, str) and value.strip().isdigit():
        n = int(value.strip())
    else:
        return None
    return n if n >= 1 else None


def _norm(value: Any) -> str:
    s = unicodedata.normalize("NFKD", "" if value is None else str(value))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


def _city_code(value: Any) -> str:
    return "" if value is None else str(value).strip().upper()


def _city_status(code: str) -> str:
    """ok · typo (ABJ mal écrit : bloquant) · other (autre ville du groupe : ignorée)."""
    if code == CITY:
        return "ok"
    if not code or _distance(code, CITY) <= 1 or sorted(code) == sorted(CITY):
        return "typo"
    return "other"


def _distance(a: str, b: str) -> int:
    previous = list(range(len(b) + 1))
    for i, ca in enumerate(a, start=1):
        current = [i]
        for j, cb in enumerate(b, start=1):
            current.append(min(previous[j] + 1, current[j - 1] + 1, previous[j - 1] + (ca != cb)))
        previous = current
    return previous[-1]


def _cell(row: Row, index: int) -> Any:
    return row[index] if index < len(row) else None


def _is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _is_empty_row(row: Row) -> bool:
    return all(_is_empty(c) for c in row)


def _raw(value: Any) -> str:
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return "" if value is None else str(value).strip()


def _header(value: Any) -> str:
    return "" if value is None else str(value).strip()


# --- Contrôles communs ----------------------------------------------------------------------


def _check_headers(result: ParseResult, header: Row, expected: list[str]) -> dict[str, int] | None:
    index: dict[str, int] = {}
    ok = True
    for i, name in enumerate(_header(c) for c in header):
        if not name:
            continue
        if name in index:
            result.error(f"Colonne {name} présente deux fois (ligne 1).", 1, name)
            ok = False
        else:
            index[name] = i
    extra = [name for name in index if name not in expected]
    for name in expected:
        if name not in index:
            message = f"Colonne {name} introuvable (ligne 1)."
            close = get_close_matches(name, extra, n=1, cutoff=0.75)
            if close:
                message += f' Colonne proche trouvée : "{close[0]}".'
            result.error(message, 1, name)
            ok = False
    cumuls = [name for name in extra if name.lower().startswith("cumul")]
    if cumuls:
        result.warn(
            f"{plural(len(cumuls), 'colonne')} de cumul ignorée{'s' if len(cumuls) > 1 else ''} "
            ": l'outil les recalcule.",
            1,
        )
    for name in extra:
        if name not in cumuls:
            result.error(f"Colonne {name} inattendue (ligne 1).", 1, name)
            ok = False
    return index if ok else None


def _ignored_message(ignored: int, scope: str, kept: int) -> str:
    verb = "seront ignorées" if ignored > 1 else "sera ignorée"
    kept_verb = "sont importées" if kept > 1 else "est importée"
    return f"{plural(ignored, 'ligne')} {scope} {verb}. {plural(kept, 'ligne')} ABJ {kept_verb}."


def _month_checks(
    result: ParseResult,
    months: list[date],
    existing: Collection[date],
    series: Mapping[str, Mapping[date, float]],
    previous: Previous,
) -> None:
    existing_set = set(existing)

    already = [m for m in months if m in existing_set]
    if len(already) == 1:
        result.warn(
            f"Le mois de {month_label(already[0])} est déjà présent : "
            "les données existantes seront remplacées."
        )
    elif already:
        result.warn(
            f"{len(already)} mois déjà présents ({month_label(already[0])} à "
            f"{month_label(already[-1])}) : les données existantes seront remplacées."
        )

    missing: list[date] = []
    for a, b in zip(months, months[1:], strict=False):
        d = add_months(a, 1)
        while d < b:
            missing.append(d)
            d = add_months(d, 1)
    if existing_set and months and max(existing_set) < months[0]:
        d = add_months(max(existing_set), 1)
        while d < months[0]:
            missing.append(d)
            d = add_months(d, 1)
    missing.sort()
    if len(missing) > 3:
        result.warn(
            f"{len(missing)} mois sans données entre {month_label(missing[0])} "
            f"et {month_label(missing[-1])}."
        )
    else:
        for d in missing:
            result.warn(f"Aucune ligne pour {month_label(d)}.")

    for name, values in series.items():
        known = {**previous.get(name, {}), **values}
        for m in months:
            before = add_months(m, -1)
            p, v = known.get(before), values.get(m)
            if p is None or v is None or p <= 0:
                continue
            change = (v - p) / p
            if abs(change) > VARIATION_THRESHOLD:
                result.warn(
                    f"{name} {CITY} : {change * 100:+.0f} % en {month_label(m)} "
                    f"par rapport à {month_label(before)}. À confirmer.",
                    column=name,
                )


# --- Fichier 1 : Réalisés commerciaux ---------------------------------------------------------


def parse_commercial(
    rows: list[Row],
    existing_months: Collection[date] = (),
    previous: Previous | None = None,
) -> ParseResult:
    result = ParseResult(kind="commercial")
    if not rows:
        result.error("Le fichier est vide.")
        return result
    index = _check_headers(result, rows[0], COMMERCIAL_HEADERS)
    if index is None:
        return result.finish()

    lines: dict[date, int] = {}
    data: dict[date, dict[str, float]] = {}
    for n, row in enumerate(rows[1:], start=2):
        if _is_empty_row(row):
            continue
        result.rows_read += 1
        if _norm(_cell(row, index["Pays"])) != COUNTRY_KEY:
            result.rows_ignored += 1
            continue
        code = _city_code(_cell(row, index["Ville"]))
        if code != CITY:
            result.error(
                f"Ligne {n} : ville {code or '(vide)'} inconnue. Seule ville acceptée : {CITY}.",
                n,
                "Ville",
            )
            continue
        month_raw = _cell(row, index["Mois"])
        month = parse_month(month_raw)
        if month is None:
            result.error(
                f'Ligne {n} : mois "{_raw(month_raw)}" non reconnu, attendu 01/07/2026.', n, "Mois"
            )
            continue
        if month in lines:
            result.error(
                f"{CITY} / {fmt_date(month)} apparaît deux fois (lignes {lines[month]} et {n}).",
                n,
                "Mois",
            )
            continue
        lines[month] = n
        values: dict[str, float] = {}
        for metric in COMMERCIAL_METRICS:
            raw = _cell(row, index[metric])
            try:
                values[metric] = parse_number(raw)
            except ValueError as exc:
                result.error(_number_message(n, metric, raw, exc), n, metric)
                continue
            if values[metric] < 0:
                result.warn(
                    f"Ligne {n}, colonne {metric} : valeur négative ({fmt_num(values[metric])}).",
                    n,
                    metric,
                )
        if len(values) == len(COMMERCIAL_METRICS):
            data[month] = values
            result.rows_kept += 1

    if result.rows_ignored:
        result.warn(_ignored_message(result.rows_ignored, "hors Côte d'Ivoire", result.rows_kept))
    if not data and result.ok:
        result.error("Aucune ligne Côte d'Ivoire / ABJ dans le fichier.")
    if not result.ok:
        return result.finish()

    months = sorted(data)
    for m in months:
        for total, parts in TOTAL_CHECKS:
            parts_sum = sum(data[m][p] for p in parts)
            if abs(data[m][total] - parts_sum) > 0.5:
                result.warn(
                    f"{CITY} {month_label(m)} : {total} {fmt_num(data[m][total])}, "
                    f"somme des offres {fmt_num(parts_sum)}.",
                    lines[m],
                    total,
                )
    series = {name: {m: data[m][name] for m in months} for name in VARIATION_METRICS}
    _month_checks(result, months, existing_months, series, previous or {})

    result.months = months
    result.records = [
        Record(CITY, m, metric, 0, value) for m in months for metric, value in data[m].items()
    ]
    return result.finish()


# --- Fichier 2 : Échus ------------------------------------------------------------------------


def _month_from_name(value: Any) -> int | None:
    if not isinstance(value, str):
        return None
    name = _norm(value)
    for prefix, number in _MONTH_PREFIXES:
        if name.startswith(prefix):
            return number
    return None


def _crosstab_months(header: Row) -> list[tuple[int, date | int]]:
    """Colonnes mois du format croisé : date complète, ou numéro de mois sans année."""
    columns: list[tuple[int, date | int]] = []
    for i, cell in enumerate(header):
        if i < 2:
            continue
        if isinstance(cell, date):
            columns.append((i, date(cell.year, cell.month, 1)))
        else:
            number = _month_from_name(cell)
            if number:
                columns.append((i, number))
    return columns


def parse_echus(
    rows: list[Row],
    year: int | None = None,
    existing_months: Collection[date] = (),
    previous: Previous | None = None,
) -> ParseResult:
    result = ParseResult(kind="echus")
    if not rows:
        result.error("Le fichier est vide.")
        return result

    header = rows[0]
    names = [_header(c) for c in header]
    # (ligne, ville, mois, ancienneté, valeur)
    entries: list[tuple[int, Any, Any, Any, Any]] = []
    if all(h in names for h in ECHUS_HEADERS):
        index = _check_headers(result, header, ECHUS_HEADERS)
        if index is None:
            return result.finish()
        for n, row in enumerate(rows[1:], start=2):
            if not _is_empty_row(row):
                entries.append(
                    (
                        n,
                        _cell(row, index["Ville"]),
                        _cell(row, index["Mois"]),
                        _cell(row, index["Anciennete_mois"]),
                        _cell(row, index["Nb_echus"]),
                    )
                )
    elif columns := _crosstab_months(header):
        needs_year = any(isinstance(m, int) for _, m in columns)
        if needs_year and year is None:
            result.error(
                "Format croisé détecté (mois en colonnes, sans année). Indiquez l'année des "
                "données pour que l'outil fasse la conversion.",
                1,
            )
            return result.finish()
        if needs_year:
            result.warn(f"Format croisé converti : mois en colonnes, année {year} appliquée.")
        for n, row in enumerate(rows[1:], start=2):
            if _is_empty_row(row):
                continue
            if _norm(_cell(row, 1)) == "villes":
                result.warn("Deuxième ligne d'en-tête détectée et ignorée.", n)
                continue
            for col, m in columns:
                value = _cell(row, col)
                if _is_empty(value):
                    continue  # mois pas encore renseigné
                cell_month = m if isinstance(m, date) else date(year or 0, m, 1)
                entries.append((n, _cell(row, 1), cell_month, _cell(row, 0), value))
    else:
        _check_headers(result, header, ECHUS_HEADERS)
        return result.finish()

    seen: dict[tuple[date, int], int] = {}
    values: dict[tuple[date, int], float] = {}
    read: set[int] = set()
    kept: set[int] = set()
    ignored: set[int] = set()
    for n, city_raw, month_raw, seniority_raw, value_raw in entries:
        code = _city_code(city_raw)
        if code == "TOTAL":
            continue
        read.add(n)
        status = _city_status(code)
        if status == "other":
            ignored.add(n)
            continue
        if status == "typo":
            result.error(
                f"Ligne {n} : ville {code or '(vide)'} inconnue. Seule ville acceptée : {CITY}.",
                n,
                "Ville",
            )
            continue
        entry_month = parse_month(month_raw)
        if entry_month is None:
            result.error(
                f'Ligne {n} : mois "{_raw(month_raw)}" non reconnu, attendu 01/07/2026.', n, "Mois"
            )
            continue
        seniority = _parse_seniority(seniority_raw)
        if seniority is None:
            result.error(
                f'Ligne {n} : ancienneté "{_raw(seniority_raw)}" invalide, '
                "attendu un entier supérieur ou égal à 1.",
                n,
                "Anciennete_mois",
            )
            continue
        try:
            value = parse_number(value_raw)
        except ValueError as exc:
            result.error(_number_message(n, "Nb_echus", value_raw, exc), n, "Nb_echus")
            continue
        if seniority > SENIORITY_WARN:
            result.warn(
                f"Ligne {n} : ancienneté {seniority} mois, inhabituelle. À confirmer.",
                n,
                "Anciennete_mois",
            )
        if value < 0:
            result.warn(f"Ligne {n} : nombre d'échus négatif ({fmt_num(value)}).", n, "Nb_echus")
        key = (entry_month, seniority)
        if key in seen:
            result.error(
                f"{CITY} / {fmt_date(entry_month)} / ancienneté {seniority} apparaît deux fois "
                f"(lignes {seen[key]} et {n}).",
                n,
                "Mois",
            )
            continue
        seen[key] = n
        values[key] = value
        kept.add(n)

    result.rows_read, result.rows_kept, result.rows_ignored = len(read), len(kept), len(ignored)
    if ignored:
        result.warn(_ignored_message(len(ignored), "d'autres villes", len(kept)))
    if not values and result.ok:
        result.error("Aucune ligne ABJ dans le fichier.")
    if not result.ok:
        return result.finish()

    months = sorted({m for m, _ in values})
    totals = dict.fromkeys(months, 0.0)
    for (m, _), value in values.items():
        totals[m] += value
    _month_checks(result, months, existing_months, {ECHUS_SERIES: totals}, previous or {})

    result.months = months
    result.records = [
        Record(CITY, m, "Nb_echus", seniority, value)
        for (m, seniority), value in sorted(values.items())
    ]
    return result.finish()


def parse_file(
    kind: str,
    rows: list[Row],
    *,
    year: int | None = None,
    existing_months: Collection[date] = (),
    previous: Previous | None = None,
) -> ParseResult:
    if kind == "commercial":
        return parse_commercial(rows, existing_months, previous)
    if kind == "echus":
        return parse_echus(rows, year, existing_months, previous)
    raise ValueError(f"Type de fichier inconnu : {kind}")
